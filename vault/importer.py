import hashlib
import io
import json
import math
from datetime import date, datetime
from uuid import uuid4

import openpyxl

from . import storage
from .models import Device

REQUIRED = {"name", "message_type", "angle_x", "angle_y", "event_time"}
NUMERIC = {"angle_x", "angle_y", "temperature", "power", "signal_intensity",
           "x_init_displacement", "y_init_displacement", "x_displacement", "y_displacement",
           "x_y_displacement", "alarm_x", "alarm_y"}


def number(value, field, row):
    if value is None or value == "":
        return None
    try:
        result = float(value)
    except (TypeError, ValueError):
        raise ValueError(f"第{row}行 {field} 不是数字") from None
    if not math.isfinite(result):
        raise ValueError(f"第{row}行 {field} 不是有限数值")
    return result


def read_excel(content):
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise ValueError("无法读取文件，请上传未损坏的 .xlsx 导出文件") from None
    try:
        if "Result" not in workbook.sheetnames:
            raise ValueError("缺少 Result 明细表；请使用现有 Excel 导出格式")
        rows = workbook["Result"].iter_rows(values_only=True)
        headers = [str(v).strip() if v is not None else "" for v in next(rows)]
        if len([v for v in headers if v]) != len(set(v for v in headers if v)):
            raise ValueError("Result 表存在重复列名")
        missing = REQUIRED - set(headers)
        if missing:
            raise ValueError("缺少必要列：" + ", ".join(sorted(missing)))
        result, skipped, warnings = [], 0, []
        for index, values in enumerate(rows, 2):
            if all(v is None for v in values):
                skipped += 1
                continue
            raw = dict(zip(headers, values))
            device = str(raw["name"]).strip() if raw["name"] is not None else ""
            if not device:
                raise ValueError(f"第{index}行缺少设备名称")
            message = str(raw["message_type"]).strip()
            if message not in {"Heartbeat", "Alarm"}:
                raise ValueError(f"第{index}行出现未知消息类型 {message}，请先确认导出口径")
            event = raw["event_time"]
            try:
                if isinstance(event, date) and not isinstance(event, datetime):
                    event = datetime.combine(event, datetime.min.time())
                if not isinstance(event, datetime):
                    event = datetime.fromisoformat(str(event).strip())
                if event.tzinfo is not None:
                    raise ValueError("带时区时间需先统一为导出本地时间")
            except (ValueError, TypeError):
                raise ValueError(f"第{index}行时间无法解析或带有时区，请确认导出时间口径") from None
            clean = {}
            for field, value in raw.items():
                if field in NUMERIC:
                    clean[field] = number(value, field, index)
                elif isinstance(value, (date, datetime)):
                    clean[field] = value.isoformat()
                else:
                    clean[field] = value
            clean.update(name=device, message_type=message,
                         event_time=event.isoformat(timespec="microseconds" if event.microsecond else "seconds"))
            for field in ["angle_x", "angle_y"]:
                if clean[field] is not None and abs(clean[field]) > 90:
                    raise ValueError(f"第{index}行 {field} 超出±90度，请确认角度单位")
            result.append(clean)
        if not result:
            raise ValueError("Result 表没有可导入记录")
        if any(r.get("x_displacement") is None or r.get("y_displacement") is None for r in result):
            warnings.append("部分记录缺少位移；保留角度分析，不将缺失位移补为零。")
        return result, skipped, warnings
    finally:
        workbook.close()


def canonical(record):
    return json.dumps(record, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False)


def import_excel(warehouse_id, filename, content):
    storage.get_warehouse(warehouse_id)
    records, blanks, warnings = read_excel(content)
    digest = hashlib.sha256(content).hexdigest()
    key = uuid4().hex
    unique, repeated = {}, 0
    for record in records:
        identity = (record["name"], record["event_time"], record["message_type"])
        payload = canonical(record)
        if identity in unique:
            if unique[identity] != payload:
                raise ValueError(f"同一文件含冲突记录：设备{identity[0]}，{identity[1]}，{identity[2]}")
            repeated += 1
        unique[identity] = payload
    with storage.database() as db:
        existing = {(r["device_id"], r["event_time"], r["message_type"]): r["payload"]
                    for r in db.execute("SELECT device_id,event_time,message_type,payload FROM records WHERE warehouse_id=?", (warehouse_id,))}
        pending = []
        for identity, payload in unique.items():
            if identity in existing:
                if existing[identity] != payload:
                    raise ValueError(f"与已导入数据冲突：设备{identity[0]}，{identity[1]}。本次未写入，请先核对源文件。")
                repeated += 1
            else:
                pending.append((warehouse_id, *identity, payload, key))
        summary = {"read": len(records), "added": len(pending), "duplicates": repeated,
                   "blank_rows": blanks, "warnings": warnings}
        source = storage.DATA / "sources" / (digest + ".xlsx")
        source.parent.mkdir(parents=True, exist_ok=True)
        if not source.exists():
            source.write_bytes(content)
        db.execute("INSERT INTO imports VALUES (?,?,?,?,?,?,?)",
                   (key, warehouse_id, filename, digest, storage.now(), canonical(summary), str(source)))
        db.executemany("INSERT INTO records VALUES (?,?,?,?,?,?)", pending)
        for device in sorted({r["name"] for r in records}):
            db.execute("INSERT OR IGNORE INTO devices VALUES (?,?,?)",
                       (warehouse_id, device, canonical(Device(id=device, label=device).model_dump())))
    return summary
